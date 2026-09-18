#!/usr/bin/env python3
# ==============================================================================
# kiw-excel - Automatic Excel (.xlsx) Exporter
# Converts Grist SQLite tables into multi-sheet Microsoft Excel workbooks
# ==============================================================================
import os
import glob
import sqlite3
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.environ.get("DATA_DIR", "/data")
BACKUP_DIR = os.environ.get("BACKUP_DIR", "/backups")
EXCEL_OUT_DIR = os.path.join(BACKUP_DIR, "excel")

os.makedirs(EXCEL_OUT_DIR, exist_ok=True)

# 1. Load document name mappings from home.sqlite3 if available
doc_names = {}
home_db = os.path.join(DATA_DIR, "home.sqlite3")
if os.path.exists(home_db):
    try:
        con = sqlite3.connect(home_db)
        cur = con.cursor()
        for doc_id, name in cur.execute("SELECT id, name FROM docs;"):
            doc_names[doc_id] = name
        con.close()
    except Exception as e:
        print(f"[Excel Export] Warning loading home.sqlite3: {e}")

# 2. Iterate through all .grist document files
grist_files = glob.glob(os.path.join(DATA_DIR, "docs", "*.grist"))
if not grist_files:
    print("[Excel Export] No .grist documents found in docs directory.")
    exit(0)

header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

for grist_path in grist_files:
    doc_id = os.path.splitext(os.path.basename(grist_path))[0]
    raw_name = doc_names.get(doc_id, doc_id)
    # Sanitize file name
    safe_name = "".join(c for c in raw_name if c.isalnum() or c in " ._-#").strip()
    out_xlsx_path = os.path.join(EXCEL_OUT_DIR, f"{safe_name}.xlsx")

    print(f"--> Exporting document '{raw_name}' ({doc_id}) to Excel: {out_xlsx_path}")

    try:
        con = sqlite3.connect(grist_path)
        cur = con.cursor()

        # Get list of user tables (skip internal Grist tables starting with _)
        tables = [
            r[0] for r in cur.execute(
                "SELECT tableId FROM _grist_Tables WHERE tableId NOT LIKE '\\_%' ESCAPE '\\' ORDER BY id;"
            ).fetchall()
        ]

        if not tables:
            print(f"    No user tables found in {doc_id}.")
            con.close()
            continue

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default empty sheet

        for table_id in tables:
            # Query column metadata for this table
            col_info = cur.execute(
                """
                SELECT colId, label, type 
                FROM _grist_Tables_column 
                WHERE parentId = (SELECT id FROM _grist_Tables WHERE tableId = ?)
                  AND colId NOT IN ('manualSort', 'id')
                ORDER BY id;
                """,
                (table_id,)
            ).fetchall()

            if not col_info:
                # Fallback to pragma if metadata table is empty
                pragma_cols = cur.execute(f'PRAGMA table_info("{table_id}")').fetchall()
                col_info = [(p[1], p[1], p[2]) for p in pragma_cols if p[1] not in ('manualSort', 'id')]

            if not col_info:
                continue

            col_ids = [c[0] for c in col_info]
            col_labels = [c[1] or c[0] for c in col_info]
            col_types = [c[2] for c in col_info]

            # Excel sheet title limited to 31 chars
            sheet_title = table_id[:31]
            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True

            # Write header row
            ws.append(col_labels)
            for col_idx in range(1, len(col_labels) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[1].height = 24

            # Fetch rows
            escaped_col_ids = ", ".join(f'"{c}"' for c in col_ids)
            rows = cur.execute(f'SELECT {escaped_col_ids} FROM "{table_id}";').fetchall()

            for row_idx, row in enumerate(rows, start=2):
                formatted_row = []
                for val, ctype in zip(row, col_types):
                    if val is None:
                        formatted_row.append("")
                    elif "Date" in ctype and isinstance(val, (int, float)) and val > 0:
                        try:
                            # Grist timestamps are in seconds
                            dt = datetime.datetime.fromtimestamp(val, tz=datetime.timezone.utc)
                            formatted_row.append(dt.strftime("%Y-%m-%d") if ctype == "Date" else dt.strftime("%Y-%m-%d %H:%M:%S"))
                        except Exception:
                            formatted_row.append(val)
                    else:
                        formatted_row.append(val)

                ws.append(formatted_row)
                for col_idx in range(1, len(formatted_row) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(out_xlsx_path)
        con.close()
        print(f"    Exported {len(tables)} tables to {out_xlsx_path} successfully.")
    except Exception as e:
        print(f"    ERROR exporting document {doc_id}: {e}")

print(f"[Excel Export] All documents exported to {EXCEL_OUT_DIR}/.")
